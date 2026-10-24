import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from analysis.pattern_analyzer import analyze_pattern
from agents import brain
from memory import context
from tools import app_tools, excel_tools, file_tools, tasks, utility_tools
from tools.learner import _validate


class ExistingFeatureRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent)
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_file_search_is_scoped_and_single_item_launch_is_routed(self):
        folder = self.root / "sandbox-files"
        folder.mkdir()
        report = folder / "report.pdf"
        report.write_text("test")
        (folder / "ignored.txt").write_text("not a pdf")
        with mock.patch.object(file_tools, "SEARCH_FOLDERS", [folder]):
            found = file_tools.find_matching_files("report", extensions=[".pdf"])
        self.assertEqual(found, [report])
        with mock.patch.object(app_tools, "subprocess") as proc:
            self.assertIn("opened", app_tools.open_calculator().lower())
            proc.Popen.assert_called_once()

    def test_excel_operations_use_only_temporary_workbook(self):
        path = self.root / "marks.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["Name", "Mark"])
        sheet.append(["Kavin", 80])
        book.save(path); book.close()
        self.assertIn("80", excel_tools.get_student_mark(path, "Kavin"))
        self.assertIn("Updated", excel_tools.update_student_mark(path, "Kavin", 95))
        self.assertIn("Added", excel_tools.add_student(path, "Madhan", 88))
        book = load_workbook(path, data_only=True)
        self.assertEqual(book.active["B2"].value, 95)
        self.assertEqual(book.active["A3"].value, "Madhan")
        book.close()

    def test_tasks_are_persisted_and_completion_updates_temporary_data(self):
        original = tasks.TASK_FILE
        tasks.TASK_FILE = self.root / "tasks.json"
        try:
            self.assertIn("Task added", tasks.add_task("Test task", "high"))
            self.assertIn("Test task", tasks.show_tasks())
            self.assertIn("completed", tasks.complete_task(1).lower())
            self.assertEqual(tasks.show_tasks(), "You have no pending tasks. ⚡")
        finally:
            tasks.TASK_FILE = original

    def test_corrupt_runtime_json_is_preserved_instead_of_overwritten(self):
        task_path = self.root / "tasks.json"
        task_path.write_text("{broken", encoding="utf-8")
        original_task_path = tasks.TASK_FILE
        tasks.TASK_FILE = task_path
        try:
            self.assertIn("preserved", tasks.add_task("Do not overwrite", "normal"))
            self.assertEqual(task_path.read_text(encoding="utf-8"), "{broken")
        finally:
            tasks.TASK_FILE = original_task_path

        memory_path = self.root / "memory.json"
        memory_path.write_text("{broken", encoding="utf-8")
        original_memory_path = context.MEMORY_FILE
        context.MEMORY_FILE = memory_path
        try:
            self.assertIn("preserved", context.remember("name", "Madhankumar"))
            self.assertEqual(memory_path.read_text(encoding="utf-8"), "{broken")
        finally:
            context.MEMORY_FILE = original_memory_path

    def test_safe_calculation_pattern_and_skill_validation(self):
        self.assertEqual(utility_tools.calculate("2 + 3 * 4"), "Answer: 14")
        self.assertIn("basic arithmetic", utility_tools.calculate("__import__('os').system('whoami')").lower())
        self.assertIn("Pattern detected:", analyze_pattern("pattern 1 2 3"))
        self.assertFalse(_validate("run_shell_command"))

    def test_unrelated_command_clears_pending_destructive_confirmations(self):
        try:
            brain.pending_file_delete = True
            self.assertEqual(brain.understand_command("open chrome")["action"], "OPEN_APP")
            self.assertFalse(brain.pending_file_delete)
            self.assertNotEqual(brain.understand_command("yes")["action"], "FILE_DELETE_CONFIRM")
            brain.pending_task_clear = True
            self.assertEqual(brain.understand_command("open calculator")["action"], "OPEN_APP")
            self.assertFalse(brain.pending_task_clear)
            self.assertNotEqual(brain.understand_command("yes")["action"], "TASK_CLEAR_ALL_CONFIRM")
        finally:
            brain.pending_file_delete = False
            brain.pending_task_clear = False

    def test_system_and_network_helpers_survive_missing_os_tools(self):
        with mock.patch.object(utility_tools.subprocess, "run", side_effect=FileNotFoundError), \
             mock.patch.object(utility_tools.socket, "gethostbyname", side_effect=OSError), \
             mock.patch.object(utility_tools.socket, "create_connection", side_effect=OSError):
            self.assertIn("Not available", utility_tools.system_information())
            report = utility_tools.network_information()
        self.assertIn("Internet: Disconnected", report)


if __name__ == "__main__":
    unittest.main()
