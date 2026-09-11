from openpyxl import load_workbook


# -------------------------------------------------
# FIND STUDENT
# -------------------------------------------------

def find_student_row(sheet, student_name, name_column):

    for row in range(2, sheet.max_row + 1):

        value = sheet.cell(
            row=row,
            column=name_column
        ).value

        if (
            value is not None
            and str(value).strip().lower()
            == student_name.strip().lower()
        ):
            return row

    return None


# -------------------------------------------------
# FIND NAME AND MARK COLUMNS
# -------------------------------------------------

def get_columns(sheet):

    name_column = None
    mark_column = None

    for cell in sheet[1]:

        if cell.value is not None:

            value = str(cell.value).strip().lower()

            if value == "name":
                name_column = cell.column

            elif value == "mark":
                mark_column = cell.column

    return name_column, mark_column


# -------------------------------------------------
# CHANGE STUDENT MARK
# -------------------------------------------------

def update_student_mark(file_path, student_name, new_mark):

    try:

        workbook = load_workbook(file_path)

        for sheet in workbook.worksheets:

            name_column, mark_column = get_columns(sheet)

            if name_column is None or mark_column is None:
                continue

            row = find_student_row(
                sheet,
                student_name,
                name_column
            )

            if row is not None:

                mark_cell = sheet.cell(
                    row=row,
                    column=mark_column
                )

                old_mark = mark_cell.value

                mark_cell.value = new_mark

                workbook.save(file_path)

                return (
                    f"Updated {student_name}'s mark "
                    f"from {old_mark} to {new_mark}."
                )

        return f"Couldn't find student '{student_name}'."

    except Exception as e:

        return f"Excel error: {e}"


# -------------------------------------------------
# GET STUDENT MARK
# -------------------------------------------------

def get_student_mark(file_path, student_name):

    try:

        workbook = load_workbook(file_path)

        for sheet in workbook.worksheets:

            name_column, mark_column = get_columns(sheet)

            if name_column is None or mark_column is None:
                continue

            row = find_student_row(
                sheet,
                student_name,
                name_column
            )

            if row is not None:

                mark = sheet.cell(
                    row=row,
                    column=mark_column
                ).value

                return f"{student_name}'s mark is {mark}."

        return f"Couldn't find student '{student_name}'."

    except Exception as e:

        return f"Excel error: {e}"


# -------------------------------------------------
# SHOW ALL STUDENTS
# -------------------------------------------------

def get_all_student_marks(file_path):

    try:

        workbook = load_workbook(file_path)

        for sheet in workbook.worksheets:

            name_column, mark_column = get_columns(sheet)

            if name_column is None or mark_column is None:
                continue

            students = []

            for row in range(2, sheet.max_row + 1):

                name = sheet.cell(
                    row=row,
                    column=name_column
                ).value

                mark = sheet.cell(
                    row=row,
                    column=mark_column
                ).value

                if name is not None:

                    students.append(
                        f"{name}: {mark}"
                    )

            if students:
                return "\n".join(students)

        return "No students found."

    except Exception as e:

        return f"Excel error: {e}"


# -------------------------------------------------
# ADD STUDENT
# -------------------------------------------------

def add_student(file_path, student_name, mark):

    try:

        workbook = load_workbook(file_path)

        for sheet in workbook.worksheets:

            name_column, mark_column = get_columns(sheet)

            if name_column is None or mark_column is None:
                continue

            existing_row = find_student_row(
                sheet,
                student_name,
                name_column
            )

            if existing_row is not None:

                return (
                    f"{student_name} already exists "
                    f"in the Excel file."
                )

            new_row = sheet.max_row + 1

            sheet.cell(
                row=new_row,
                column=name_column
            ).value = student_name

            sheet.cell(
                row=new_row,
                column=mark_column
            ).value = mark

            workbook.save(file_path)

            return (
                f"Added {student_name} "
                f"with mark {mark}."
            )

        return "Could not find NAME and MARK columns."

    except Exception as e:

        return f"Excel error: {e}"


# -------------------------------------------------
# INCREASE MARK
# -------------------------------------------------

def increase_student_mark(file_path, student_name, amount):

    try:

        workbook = load_workbook(file_path)

        for sheet in workbook.worksheets:

            name_column, mark_column = get_columns(sheet)

            if name_column is None or mark_column is None:
                continue

            row = find_student_row(
                sheet,
                student_name,
                name_column
            )

            if row is not None:

                mark_cell = sheet.cell(
                    row=row,
                    column=mark_column
                )

                old_mark = mark_cell.value

                new_mark = old_mark + amount

                mark_cell.value = new_mark

                workbook.save(file_path)

                return (
                    f"Increased {student_name}'s mark "
                    f"from {old_mark} to {new_mark}."
                )

        return f"Couldn't find student '{student_name}'."

    except Exception as e:

        return f"Excel error: {e}"
