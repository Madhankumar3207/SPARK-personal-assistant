from pathlib import Path
from openpyxl import load_workbook

file_path = Path("test_files/test.xlsx").resolve()

print("FILE BEING USED:")
print(file_path)

print("\nBEFORE:")
workbook = load_workbook(file_path)
sheet = workbook.active

print("B3 =", sheet["B3"].value)

print("\nCHANGING B3 TO 95...")

sheet["B3"] = 95

workbook.save(file_path)

print("Saved!")

print("\nCHECKING FILE AGAIN:")

workbook.close()

workbook = load_workbook(file_path, data_only=False)
sheet = workbook.active

print("B3 =", sheet["B3"].value)

workbook.close()